"""Point 4(b): Bonanos+ 2009 (J/AJ/138/1003, table 3) -> CSV, Excel with chart, CMD PNG, VOTable.

Source: machine-readable table 3 of the AJ article (VizieR was unreachable on 2026-10-28).
Run: .venv/bin/python resultados/scripts/p4b_bonanos_cmd.py
"""

from pathlib import Path

import aladin_mcp  # noqa: F401  (keeps caches inside the project)
import matplotlib.pyplot as plt
import numpy as np
from astropy.io import ascii
from astropy.table import Table
from openpyxl import Workbook
from openpyxl.chart import Reference, ScatterChart, Series

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "resultados" / "punto4"
SRC = OUT / "bonanos2009_original" / "aj318206t3_mrt.txt"

table = ascii.read(SRC, format="mrt")
df = table.to_pandas()
df["B_V"] = (df["Bmag"] - df["Vmag"]).round(3)
df["clase"] = df["CandC"].str.strip().str.extract(r"^(WN|WC|WO|LBV|O|B|A|F|G|K|M)", expand=False)
df.loc[df["clase"].isin(["WN", "WC", "WO"]), "clase"] = "WR"
cmd = df.dropna(subset=["Vmag", "B_V"]).copy()
# Intrinsic B-V of the hottest stars is about -0.33; values far outside [-0.5, 2.5]
# point to B and V from different sources/epochs or blends, not to real colours.
cmd["BV_plausible"] = cmd["B_V"].between(-0.5, 2.5)

keep = ["Name", "RAdeg", "DEdeg", "Bmag", "Vmag", "B_V", "BV_plausible", "clase", "CandC"]
cmd[keep].to_csv(OUT / "Bonanos2009_LMC_V_BV.csv", index=False)
df.to_csv(OUT / "Bonanos2009_LMC_tabla3_completa.csv", index=False)
Table.from_pandas(cmd[keep]).write(OUT / "Bonanos2009_LMC_para_Aladin.xml", format="votable", overwrite=True)

# ---- Excel: data sheet + native scatter chart with the V axis reversed
wb = Workbook()
ws = wb.active
ws.title = "datos"
ws.append(keep)
good = cmd[cmd["BV_plausible"]]
for row in good[keep].itertuples(index=False):
    ws.append([None if (isinstance(v, float) and np.isnan(v)) else v for v in row])
n = len(good) + 1
wa = wb.create_sheet("descartadas_BV_implausible")
wa.append(keep)
for row in cmd[~cmd["BV_plausible"]][keep].itertuples(index=False):
    wa.append([None if (isinstance(v, float) and np.isnan(v)) else v for v in row])
chart = ScatterChart()
chart.title = "LMC massive stars (Bonanos+ 2009): V vs B-V"
chart.style = 13
chart.x_axis.title = "B - V (mag)"
chart.y_axis.title = "V (mag)"
chart.y_axis.scaling.orientation = "maxMin"  # brighter stars at the top
chart.x_axis.delete = False
chart.y_axis.delete = False
series = Series(Reference(ws, min_col=5, min_row=2, max_row=n),
                Reference(ws, min_col=6, min_row=2, max_row=n), title="stars")
series.marker.symbol = "circle"
series.marker.size = 3
series.graphicalProperties.line.noFill = True
chart.series.append(series)
chart.height, chart.width = 15, 20
wc = wb.create_sheet("grafico")
wc.add_chart(chart, "A1")
wb.save(OUT / "Bonanos2009_LMC_V_BV.xlsx")

# ---- CMD with hot/cool zones
colors = {"O": "#3b5bdb", "WR": "#7048e8", "B": "#4dabf7", "LBV": "#e64980", "A": "#adb5bd",
          "F": "#fab005", "G": "#fd7e14", "K": "#f76707", "M": "#c92a2a"}
fig, ax = plt.subplots(figsize=(8, 8))
bad = cmd[~cmd["BV_plausible"]]
ax.scatter(bad["B_V"].clip(-0.75, 2.75), bad["Vmag"], s=14, marker="x", color="0.6",
           label=f"implausible B-V ({len(bad)}), clipped")
for cls, color in colors.items():
    sel = good[good["clase"] == cls]
    if len(sel):
        ax.scatter(sel["B_V"], sel["Vmag"], s=9, color=color, label=f"{cls} ({len(sel)})", alpha=0.8)
other = good[good["clase"].isna()]
if len(other):
    ax.scatter(other["B_V"], other["Vmag"], s=9, color="k", label=f"no class ({len(other)})")
ax.invert_yaxis()
ax.set_xlim(-0.8, 2.8)
ax.axvspan(-0.8, 0.0, color="#4dabf7", alpha=0.08)
ax.axvspan(1.2, 2.8, color="#c92a2a", alpha=0.08)
ax.text(0.03, 0.04, "HOTTEST\n(B-V < 0: blue)\nO, B, WR stars", transform=ax.transAxes,
        color="#1c3faa", fontsize=10, va="bottom")
ax.text(0.97, 0.04, "COOLEST\n(B-V > 1.2: red)\nK, M supergiants", transform=ax.transAxes,
        color="#a51111", fontsize=10, va="bottom", ha="right")
ax.set_xlabel("Colour index B - V (mag)")
ax.set_ylabel("Apparent visual magnitude V (mag)")
ax.set_title(f"Massive stars in the LMC - Bonanos et al. 2009 (N = {len(good)} + {len(bad)} flagged)")
ax.grid(alpha=0.3)
ax.legend(loc="upper right", fontsize=8, title="spectral class")
fig.savefig(OUT / "Bonanos2009_LMC_V_vs_BV.png", dpi=130, bbox_inches="tight")

summary = good.groupby("clase")["B_V"].describe()[["count", "mean", "min", "max"]].round(2)
summary.to_csv(OUT / "Bonanos2009_resumen_por_clase.csv")
print(summary.to_string())
print("plausible:", len(good), "flagged:", len(bad))
print("V range:", good["Vmag"].min(), good["Vmag"].max(), "| B-V range:", good["B_V"].min(), good["B_V"].max())
